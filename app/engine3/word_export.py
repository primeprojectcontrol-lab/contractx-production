from __future__ import annotations

import io
import re

from datetime import datetime, timezone
from typing import Optional

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# =====================================================================
# BRAND / DOCUMENT CONSTANTS
# =====================================================================

BRAND_NAME = "contractX.co"
BRAND_TAGLINE = "AI Contract & Claims Xpert for Construction"

COLOR_NAVY = "1F2937"
COLOR_STEEL = "475569"
COLOR_GOLD = "B99146"
COLOR_LIGHT_GREY = "F3F4F6"
COLOR_BORDER = "D1D5DB"
COLOR_WHITE = "FFFFFF"
COLOR_RED = "C00000"
COLOR_BLACK = "111827"

DISCLAIMER_HEADING = "AI SUPPORT DISCLAIMER"

DISCLAIMER_TEXT = (
    "This document is AI-assisted decision-support material only. "
    "It may contain errors, omissions or incorrect interpretations. "
    "It does not constitute legal advice, contractual certification, "
    "expert opinion or a binding determination. All contractual, legal, "
    "technical and commercial conclusions must be independently reviewed "
    "against the executed contract, project records and applicable requirements. "
    "Final review, contextual interpretation, reliance and decision-making "
    "remain solely the responsibility of the authorised human reviewer."
)


# =====================================================================
# HELPERS
# =====================================================================

def _safe_filename(value: str) -> str:
    value = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        value or "",
    )

    return (
        value.strip("_")
        or "ContractX_Draft"
    )


def _set_cell_shading(
    cell,
    fill: str,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()

    shd = tc_pr.find(
        qn("w:shd")
    )

    if shd is None:
        shd = OxmlElement(
            "w:shd"
        )
        tc_pr.append(
            shd
        )

    shd.set(
        qn("w:fill"),
        fill,
    )


def _set_cell_border(
    cell,
    *,
    color: str = COLOR_BORDER,
    size: str = "4",
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()

    tc_borders = tc_pr.first_child_found_in(
        "w:tcBorders"
    )

    if tc_borders is None:
        tc_borders = OxmlElement(
            "w:tcBorders"
        )
        tc_pr.append(
            tc_borders
        )

    for edge in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    ):
        tag = f"w:{edge}"

        edge_el = tc_borders.find(
            qn(tag)
        )

        if edge_el is None:
            edge_el = OxmlElement(
                tag
            )
            tc_borders.append(
                edge_el
            )

        edge_el.set(
            qn("w:val"),
            "single",
        )

        edge_el.set(
            qn("w:sz"),
            size,
        )

        edge_el.set(
            qn("w:space"),
            "0",
        )

        edge_el.set(
            qn("w:color"),
            color,
        )


def _add_horizontal_rule(
    paragraph,
    *,
    color: str = COLOR_GOLD,
    size: str = "8",
) -> None:
    p = paragraph._p
    p_pr = p.get_or_add_pPr()

    p_bdr = p_pr.find(
        qn("w:pBdr")
    )

    if p_bdr is None:
        p_bdr = OxmlElement(
            "w:pBdr"
        )
        p_pr.append(
            p_bdr
        )

    bottom = OxmlElement(
        "w:bottom"
    )

    bottom.set(
        qn("w:val"),
        "single",
    )

    bottom.set(
        qn("w:sz"),
        size,
    )

    bottom.set(
        qn("w:space"),
        "1",
    )

    bottom.set(
        qn("w:color"),
        color,
    )

    p_bdr.append(
        bottom
    )


def _add_page_number(
    paragraph,
) -> None:
    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    run = paragraph.add_run(
        "Page "
    )

    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(
        COLOR_STEEL
    )

    fld_char_begin = OxmlElement(
        "w:fldChar"
    )
    fld_char_begin.set(
        qn("w:fldCharType"),
        "begin",
    )

    instr_text = OxmlElement(
        "w:instrText"
    )
    instr_text.set(
        qn("xml:space"),
        "preserve",
    )
    instr_text.text = "PAGE"

    fld_char_end = OxmlElement(
        "w:fldChar"
    )
    fld_char_end.set(
        qn("w:fldCharType"),
        "end",
    )

    run._r.append(
        fld_char_begin
    )
    run._r.append(
        instr_text
    )
    run._r.append(
        fld_char_end
    )


def _apply_run_font(
    run,
    *,
    size: float = 10.5,
    bold: Optional[bool] = None,
    italic: Optional[bool] = None,
    color: str = COLOR_BLACK,
) -> None:
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(
        color
    )

    if bold is not None:
        run.bold = bold

    if italic is not None:
        run.italic = italic


def _add_markdown_runs(
    paragraph,
    text: str,
    *,
    font_size: float = 10.5,
    color: str = COLOR_BLACK,
) -> None:
    """
    Supports simple inline markdown:
    **bold**
    *italic*
    """

    pattern = re.compile(
        r"(\*\*.+?\*\*|\*.+?\*)"
    )

    cursor = 0

    for match in pattern.finditer(
        text
    ):
        if match.start() > cursor:
            run = paragraph.add_run(
                text[cursor:match.start()]
            )
            _apply_run_font(
                run,
                size=font_size,
                color=color,
            )

        token = match.group(0)

        if token.startswith(
            "**"
        ) and token.endswith(
            "**"
        ):
            run = paragraph.add_run(
                token[2:-2]
            )
            _apply_run_font(
                run,
                size=font_size,
                bold=True,
                color=color,
            )

        elif token.startswith(
            "*"
        ) and token.endswith(
            "*"
        ):
            run = paragraph.add_run(
                token[1:-1]
            )
            _apply_run_font(
                run,
                size=font_size,
                italic=True,
                color=color,
            )

        cursor = match.end()

    if cursor < len(
        text
    ):
        run = paragraph.add_run(
            text[cursor:]
        )
        _apply_run_font(
            run,
            size=font_size,
            color=color,
        )


def _configure_document(
    doc: Document,
) -> None:
    section = doc.sections[0]

    section.top_margin = Inches(
        0.72
    )
    section.bottom_margin = Inches(
        0.78
    )
    section.left_margin = Inches(
        0.8
    )
    section.right_margin = Inches(
        0.8
    )

    section.header_distance = Inches(
        0.28
    )
    section.footer_distance = Inches(
        0.30
    )

    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(
        COLOR_BLACK
    )

    normal.paragraph_format.space_after = Pt(
        6
    )
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in (
        (
            "Title",
            20,
            COLOR_NAVY,
        ),
        (
            "Heading 1",
            15,
            COLOR_NAVY,
        ),
        (
            "Heading 2",
            12.5,
            COLOR_NAVY,
        ),
        (
            "Heading 3",
            11,
            COLOR_STEEL,
        ),
    ):
        style = styles[
            style_name
        ]

        style.font.name = "Aptos Display"
        style.font.size = Pt(
            size
        )
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(
            color
        )

        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(
            9
        )
        style.paragraph_format.space_after = Pt(
            4
        )


def _build_header(
    doc: Document,
) -> None:
    for section in doc.sections:
        header = section.header

        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        brand = p.add_run(
            "contract"
        )
        _apply_run_font(
            brand,
            size=10,
            bold=True,
            color=COLOR_NAVY,
        )

        x_run = p.add_run(
            "X"
        )
        _apply_run_font(
            x_run,
            size=10,
            bold=True,
            color=COLOR_GOLD,
        )

        tail = p.add_run(
            ".co"
        )
        _apply_run_font(
            tail,
            size=10,
            bold=True,
            color=COLOR_NAVY,
        )

        tagline = p.add_run(
            f"   |   {BRAND_TAGLINE}"
        )
        _apply_run_font(
            tagline,
            size=7.5,
            color=COLOR_STEEL,
        )

        _add_horizontal_rule(
            p,
            color=COLOR_GOLD,
            size="6",
        )


def _build_footer(
    doc: Document,
) -> None:
    for section in doc.sections:
        footer = section.footer

        # Remove default footer text while retaining the paragraph.
        first = footer.paragraphs[0]
        first.clear()

        disclaimer = first
        disclaimer.alignment = (
            WD_ALIGN_PARAGRAPH.JUSTIFY
        )

        heading = disclaimer.add_run(
            f"{DISCLAIMER_HEADING}: "
        )
        _apply_run_font(
            heading,
            size=7.5,
            bold=True,
            color=COLOR_RED,
        )

        body = disclaimer.add_run(
            DISCLAIMER_TEXT
        )
        _apply_run_font(
            body,
            size=7.5,
            color=COLOR_RED,
        )

        disclaimer.paragraph_format.space_after = Pt(
            2
        )
        disclaimer.paragraph_format.line_spacing = 0.9

        page = footer.add_paragraph()
        _add_page_number(
            page
        )

        page.paragraph_format.space_before = Pt(
            1
        )
        page.paragraph_format.space_after = Pt(
            0
        )


def _add_metadata_table(
    doc: Document,
    *,
    project_id: str,
    document_type: str,
    sender_role: Optional[str],
    recipient_role: Optional[str],
    matter_reference: Optional[str],
) -> None:
    rows = [
        (
            "Project ID",
            str(project_id),
        ),
        (
            "Document Type",
            document_type,
        ),
    ]

    if matter_reference:
        rows.append(
            (
                "Matter Reference",
                matter_reference,
            )
        )

    if sender_role:
        rows.append(
            (
                "From",
                sender_role,
            )
        )

    if recipient_role:
        rows.append(
            (
                "To",
                recipient_role,
            )
        )

    rows.append(
        (
            "Generated",
            datetime.now(
                timezone.utc
            ).strftime(
                "%d %b %Y %H:%M UTC"
            ),
        )
    )

    table = doc.add_table(
        rows=len(rows),
        cols=2,
    )

    table.alignment = (
        WD_TABLE_ALIGNMENT.LEFT
    )

    table.autofit = True

    for idx, (
        label,
        value,
    ) in enumerate(
        rows
    ):
        label_cell = table.cell(
            idx,
            0,
        )

        value_cell = table.cell(
            idx,
            1,
        )

        _set_cell_shading(
            label_cell,
            COLOR_LIGHT_GREY,
        )

        for cell in (
            label_cell,
            value_cell,
        ):
            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            _set_cell_border(
                cell
            )

        label_paragraph = label_cell.paragraphs[0]
        label_paragraph.paragraph_format.space_after = Pt(
            0
        )

        run = label_paragraph.add_run(
            label
        )
        _apply_run_font(
            run,
            size=8.5,
            bold=True,
            color=COLOR_STEEL,
        )

        value_paragraph = value_cell.paragraphs[0]
        value_paragraph.paragraph_format.space_after = Pt(
            0
        )

        run = value_paragraph.add_run(
            value or ""
        )
        _apply_run_font(
            run,
            size=8.5,
            color=COLOR_BLACK,
        )

    doc.add_paragraph("")


def _add_body_content(
    doc: Document,
    content: str,
) -> None:
    """
    Converts common markdown-like AI output into professional DOCX structure.

    Supports:
    # Heading
    ## Heading
    ### Heading
    - bullet
    * bullet
    1. numbered item
    ---
    plain paragraphs
    """

    lines = content.splitlines()

    pending_plain: list[str] = []

    def flush_plain() -> None:
        if not pending_plain:
            return

        combined = " ".join(
            part.strip()
            for part in pending_plain
            if part.strip()
        ).strip()

        pending_plain.clear()

        if not combined:
            return

        p = doc.add_paragraph()
        p.alignment = (
            WD_ALIGN_PARAGRAPH.JUSTIFY
        )

        p.paragraph_format.space_after = Pt(
            6
        )

        _add_markdown_runs(
            p,
            combined,
            font_size=10.5,
        )

    for raw_line in lines:
        line = raw_line.strip()

        if not line:
            flush_plain()
            continue

        if line in (
            "---",
            "***",
        ):
            flush_plain()
            rule = doc.add_paragraph()
            _add_horizontal_rule(
                rule,
                color=COLOR_BORDER,
                size="4",
            )
            continue

        if line.startswith(
            "### "
        ):
            flush_plain()

            text = line[4:].strip()

            heading = doc.add_heading(
                "",
                level=3,
            )

            _add_markdown_runs(
                heading,
                text,
                font_size=11,
                color=COLOR_STEEL,
            )

            continue

        if line.startswith(
            "## "
        ):
            flush_plain()

            text = line[3:].strip()

            heading = doc.add_heading(
                "",
                level=2,
            )

            _add_markdown_runs(
                heading,
                text,
                font_size=12.5,
                color=COLOR_NAVY,
            )

            continue

        if line.startswith(
            "# "
        ):
            flush_plain()

            text = line[2:].strip()

            heading = doc.add_heading(
                "",
                level=1,
            )

            _add_markdown_runs(
                heading,
                text,
                font_size=15,
                color=COLOR_NAVY,
            )

            continue

        if re.match(
            r"^[-*]\s+",
            line,
        ):
            flush_plain()

            text = re.sub(
                r"^[-*]\s+",
                "",
                line,
            )

            p = doc.add_paragraph(
                style="List Bullet"
            )

            _add_markdown_runs(
                p,
                text,
                font_size=10.5,
            )

            p.paragraph_format.space_after = Pt(
                3
            )

            continue

        if re.match(
            r"^\d+\.\s+",
            line,
        ):
            flush_plain()

            text = re.sub(
                r"^\d+\.\s+",
                "",
                line,
            )

            p = doc.add_paragraph(
                style="List Number"
            )

            _add_markdown_runs(
                p,
                text,
                font_size=10.5,
            )

            p.paragraph_format.space_after = Pt(
                3
            )

            continue

        pending_plain.append(
            line
        )

    flush_plain()


# =====================================================================
# WORD EXPORTER
# =====================================================================

def build_word_document(
    *,
    title: str,
    content: str,
    project_id: str,
    document_type: str,
    sender_role: Optional[str] = None,
    recipient_role: Optional[str] = None,
    matter_reference: Optional[str] = None,
) -> tuple[io.BytesIO, str]:

    buffer = io.BytesIO()

    doc = Document()

    # -------------------------------------------------------------
    # DOCUMENT SETUP
    # -------------------------------------------------------------

    _configure_document(
        doc
    )

    _build_header(
        doc
    )

    _build_footer(
        doc
    )

    # -------------------------------------------------------------
    # TITLE BLOCK
    # -------------------------------------------------------------

    title_paragraph = doc.add_paragraph()

    title_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )

    title_paragraph.paragraph_format.space_before = Pt(
        8
    )

    title_paragraph.paragraph_format.space_after = Pt(
        4
    )

    title_run = title_paragraph.add_run(
        title
    )

    _apply_run_font(
        title_run,
        size=20,
        bold=True,
        color=COLOR_NAVY,
    )

    subtitle = doc.add_paragraph()

    subtitle.paragraph_format.space_after = Pt(
        8
    )

    subtitle_run = subtitle.add_run(
        "Professional Contractual Draft | For Human Review"
    )

    _apply_run_font(
        subtitle_run,
        size=9,
        italic=True,
        color=COLOR_STEEL,
    )

    accent = doc.add_paragraph()
    _add_horizontal_rule(
        accent,
        color=COLOR_GOLD,
        size="12",
    )

    accent.paragraph_format.space_after = Pt(
        8
    )

    # -------------------------------------------------------------
    # DOCUMENT METADATA
    # -------------------------------------------------------------

    _add_metadata_table(
        doc,
        project_id=project_id,
        document_type=document_type,
        sender_role=sender_role,
        recipient_role=recipient_role,
        matter_reference=matter_reference,
    )

    # -------------------------------------------------------------
    # GENERATED CONTENT
    # -------------------------------------------------------------

    _add_body_content(
        doc,
        content,
    )

    # -------------------------------------------------------------
    # CLOSING NOTE
    # -------------------------------------------------------------

    closing = doc.add_paragraph()

    closing.paragraph_format.space_before = Pt(
        12
    )

    _add_horizontal_rule(
        closing,
        color=COLOR_BORDER,
        size="4",
    )

    note = doc.add_paragraph()

    note.alignment = (
        WD_ALIGN_PARAGRAPH.LEFT
    )

    note_run = note.add_run(
        "Prepared using contractX.co | Review. Monitor. Resolve."
    )

    _apply_run_font(
        note_run,
        size=8.5,
        italic=True,
        color=COLOR_STEEL,
    )

    # -------------------------------------------------------------
    # CORE PROPERTIES
    # -------------------------------------------------------------

    doc.core_properties.title = title
    doc.core_properties.subject = document_type
    doc.core_properties.author = BRAND_NAME
    doc.core_properties.comments = (
        "AI-assisted draft prepared for authorised human review."
    )

    # -------------------------------------------------------------
    # SAVE
    # -------------------------------------------------------------

    doc.save(
        buffer
    )

    buffer.seek(
        0
    )

    filename = (
        f"ContractX_"
        f"{_safe_filename(document_type)}_"
        f"{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}"
        f".docx"
    )

    return (
        buffer,
        filename,
    )
