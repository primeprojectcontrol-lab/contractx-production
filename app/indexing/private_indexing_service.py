from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import os
import re
import uuid
from datetime import datetime, timezone
from typing import Any, Optional

from app.connectors.connector_service import ConnectorService


SUPABASE_URL = os.getenv(
    "SUPABASE_URL",
    "",
).strip()

SUPABASE_SERVICE_ROLE_KEY = os.getenv(
    "SUPABASE_SERVICE_ROLE_KEY",
    "",
).strip()

SUPABASE_KEY = os.getenv(
    "SUPABASE_KEY",
    "",
).strip()

BACKEND_SUPABASE_KEY = (
    SUPABASE_SERVICE_ROLE_KEY
    or SUPABASE_KEY
)

OPENAI_API_KEY = os.getenv(
    "OPENAI_API_KEY",
    "",
).strip()

PROJECT_EMBEDDING_MODEL = os.getenv(
    "PROJECT_EMBEDDING_MODEL",
    "",
).strip() or os.getenv(
    "QDRANT_EMBEDDING_MODEL",
    "",
).strip() or "text-embedding-3-large"

PROJECT_QDRANT_COLLECTION = os.getenv(
    "PROJECT_QDRANT_COLLECTION",
    "contractx_project_evidence",
).strip()

QDRANT_URL = os.getenv(
    "QDRANT_URL",
    "",
).strip()

QDRANT_HOST = os.getenv(
    "QDRANT_HOST",
    "localhost",
).strip()

QDRANT_PORT = int(
    os.getenv(
        "QDRANT_PORT",
        "6333",
    )
)

QDRANT_API_KEY = os.getenv(
    "QDRANT_API_KEY",
    "",
).strip()

PROJECT_CHUNK_CHARS = int(
    os.getenv(
        "PROJECT_CHUNK_CHARS",
        "5000",
    )
)

PROJECT_CHUNK_OVERLAP = int(
    os.getenv(
        "PROJECT_CHUNK_OVERLAP",
        "600",
    )
)

PROJECT_EMBED_BATCH = int(
    os.getenv(
        "PROJECT_EMBED_BATCH",
        "32",
    )
)

PROJECT_DOCUMENT_BUCKET = os.getenv(
    "PROJECT_DOCUMENT_BUCKET",
    "project-documents",
).strip()


def _utc_now() -> str:
    return datetime.now(
        timezone.utc
    ).isoformat()


def _normalise_whitespace(
    text: str,
) -> str:
    text = text.replace(
        "\x00",
        " ",
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()


def _extract_pdf(
    raw: bytes,
) -> str:
    from pypdf import PdfReader

    reader = PdfReader(
        io.BytesIO(raw)
    )

    pages: list[str] = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        extracted = (
            page.extract_text()
            or ""
        ).strip()

        if extracted:
            pages.append(
                (
                    f"[Page {page_number}]\n"
                    f"{extracted}"
                )
            )

    return "\n\n".join(
        pages
    )


def _extract_docx(
    raw: bytes,
) -> str:
    from docx import Document

    doc = Document(
        io.BytesIO(raw)
    )

    blocks: list[str] = []

    for paragraph in doc.paragraphs:
        value = paragraph.text.strip()

        if value:
            blocks.append(
                value
            )

    for table in doc.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(values):
                blocks.append(
                    " | ".join(values)
                )

    return "\n".join(
        blocks
    )


def _extract_xlsx(
    raw: bytes,
) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(
        io.BytesIO(raw),
        read_only=True,
        data_only=True,
    )

    output: list[str] = []

    try:
        for sheet in workbook.worksheets:
            output.append(
                f"[Sheet: {sheet.title}]"
            )

            for row in sheet.iter_rows(
                values_only=True
            ):
                values = [
                    ""
                    if value is None
                    else str(value)
                    for value in row
                ]

                if any(
                    value.strip()
                    for value in values
                ):
                    output.append(
                        " | ".join(
                            values
                        )
                    )
    finally:
        workbook.close()

    return "\n".join(
        output
    )


def _extract_pptx(
    raw: bytes,
) -> str:
    from pptx import Presentation

    presentation = Presentation(
        io.BytesIO(raw)
    )

    output: list[str] = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_text: list[str] = []

        for shape in slide.shapes:
            value = getattr(
                shape,
                "text",
                "",
            )

            if value and value.strip():
                slide_text.append(
                    value.strip()
                )

        if slide_text:
            output.append(
                (
                    f"[Slide {slide_number}]\n"
                    + "\n".join(
                        slide_text
                    )
                )
            )

    return "\n\n".join(
        output
    )


def _extract_csv(
    raw: bytes,
) -> str:
    text = raw.decode(
        "utf-8-sig",
        errors="ignore",
    )

    reader = csv.reader(
        io.StringIO(text)
    )

    return "\n".join(
        " | ".join(row)
        for row in reader
        if any(
            value.strip()
            for value in row
        )
    )


def extract_text(
    raw: bytes,
    file_name: str,
) -> str:
    extension = (
        file_name.rsplit(
            ".",
            1,
        )[-1].lower()
        if "." in file_name
        else ""
    )

    if extension == "pdf":
        text = _extract_pdf(
            raw
        )

    elif extension == "docx":
        text = _extract_docx(
            raw
        )

    elif extension == "xlsx":
        text = _extract_xlsx(
            raw
        )

    elif extension == "pptx":
        text = _extract_pptx(
            raw
        )

    elif extension == "csv":
        text = _extract_csv(
            raw
        )

    elif extension in {
        "txt",
        "md",
        "xml",
        "json",
        "html",
        "htm",
        "eml",
        "log",
    }:
        text = raw.decode(
            "utf-8",
            errors="ignore",
        )

    else:
        raise RuntimeError(
            (
                f"Indexing is not implemented "
                f"for '.{extension or 'unknown'}' files. "
                "Use PDF, DOCX, XLSX, PPTX, CSV, TXT, "
                "XML, JSON, HTML or EML."
            )
        )

    text = _normalise_whitespace(
        text
    )

    if not text:
        raise RuntimeError(
            (
                "No machine-readable text could be extracted. "
                "Scanned-image OCR is not enabled in this pipeline."
            )
        )

    return text


def chunk_text(
    text: str,
) -> list[str]:
    chunk_size = max(
        PROJECT_CHUNK_CHARS,
        1000,
    )

    overlap = max(
        min(
            PROJECT_CHUNK_OVERLAP,
            chunk_size // 3,
        ),
        0,
    )

    paragraphs = [
        value.strip()
        for value in re.split(
            r"\n\s*\n",
            text,
        )
        if value.strip()
    ]

    chunks: list[str] = []
    current = ""

    for paragraph in paragraphs:
        candidate = (
            paragraph
            if not current
            else current
            + "\n\n"
            + paragraph
        )

        if (
            len(candidate)
            <= chunk_size
        ):
            current = candidate
            continue

        if current:
            chunks.append(
                current.strip()
            )

        if (
            len(paragraph)
            <= chunk_size
        ):
            prefix = (
                current[-overlap:]
                if (
                    overlap
                    and current
                )
                else ""
            )

            current = (
                (
                    prefix
                    + "\n\n"
                    + paragraph
                ).strip()
            )

        else:
            start = 0

            while (
                start
                < len(paragraph)
            ):
                end = min(
                    start + chunk_size,
                    len(paragraph),
                )

                chunks.append(
                    paragraph[
                        start:end
                    ].strip()
                )

                if end >= len(
                    paragraph
                ):
                    current = ""
                    break

                start = max(
                    end - overlap,
                    start + 1,
                )

    if current.strip():
        chunks.append(
            current.strip()
        )

    return [
        chunk
        for chunk in chunks
        if chunk
    ]


class PrivateProjectIndexer:
    """
    ContractX private-project ingestion/indexing.

    Reference library:
        standard_contract_library

    Private project evidence:
        contractx_project_evidence

    They are deliberately separate collections.
    """

    def __init__(
        self,
    ) -> None:
        if (
            not SUPABASE_URL
            or not BACKEND_SUPABASE_KEY
        ):
            raise RuntimeError(
                "Supabase backend configuration is missing."
            )

        try:
            from supabase import create_client

            self.db = create_client(
                SUPABASE_URL,
                BACKEND_SUPABASE_KEY,
            )

        except Exception as exc:
            raise RuntimeError(
                "Unable to initialize Supabase."
            ) from exc

        if not OPENAI_API_KEY:
            raise RuntimeError(
                (
                    "OPENAI_API_KEY is required "
                    "for private project embeddings."
                )
            )

        try:
            from openai import OpenAI

            self.openai = OpenAI(
                api_key=OPENAI_API_KEY,
                timeout=60.0,
            )

        except Exception as exc:
            raise RuntimeError(
                "Unable to initialize the OpenAI embeddings client."
            ) from exc

        try:
            from qdrant_client import QdrantClient

            if QDRANT_URL:
                self.qdrant = QdrantClient(
                    url=QDRANT_URL,
                    api_key=(
                        QDRANT_API_KEY
                        or None
                    ),
                    timeout=60,
                )
            else:
                self.qdrant = QdrantClient(
                    host=QDRANT_HOST,
                    port=QDRANT_PORT,
                    api_key=(
                        QDRANT_API_KEY
                        or None
                    ),
                    timeout=60,
                )

        except Exception as exc:
            raise RuntimeError(
                "Unable to initialize Qdrant."
            ) from exc

    def _update_document(
        self,
        document_id: str,
        values: dict[str, Any],
    ) -> None:
        values = {
            **values,
            "updated_at":
                _utc_now(),
        }

        (
            self.db.table(
                "project_documents"
            )
            .update(
                values
            )
            .eq(
                "id",
                document_id,
            )
            .execute()
        )

    def _update_job(
        self,
        job_id: Optional[str],
        values: dict[str, Any],
    ) -> None:
        if not job_id:
            return

        values = {
            **values,
            "updated_at":
                _utc_now(),
        }

        (
            self.db.table(
                "document_ingestion_jobs"
            )
            .update(
                values
            )
            .eq(
                "id",
                job_id,
            )
            .execute()
        )

    def _get_document(
        self,
        *,
        document_id: str,
        tenant_id: str,
        project_id: str,
    ) -> dict[str, Any]:
        result = (
            self.db.table(
                "project_documents"
            )
            .select("*")
            .eq(
                "id",
                document_id,
            )
            .eq(
                "tenant_id",
                tenant_id,
            )
            .eq(
                "project_id",
                project_id,
            )
            .limit(1)
            .execute()
        )

        if not result.data:
            raise RuntimeError(
                "Project document was not found."
            )

        return result.data[0]

    async def _load_bytes(
        self,
        document: dict[str, Any],
    ) -> bytes:
        storage_mode = str(
            document.get(
                "storage_mode"
            )
            or ""
        ).lower()

        if storage_mode == "contractx_managed":
            storage_path = str(
                document.get(
                    "storage_path"
                )
                or ""
            )

            if not storage_path:
                raise RuntimeError(
                    "ContractX Vault storage path is missing."
                )

            def download() -> bytes:
                value = (
                    self.db.storage.from_(
                        PROJECT_DOCUMENT_BUCKET
                    )
                    .download(
                        storage_path
                    )
                )

                if isinstance(
                    value,
                    bytes,
                ):
                    return value

                if hasattr(
                    value,
                    "content",
                ):
                    return bytes(
                        value.content
                    )

                return bytes(
                    value
                )

            return await asyncio.to_thread(
                download
            )

        source_id = str(
            document.get(
                "source_id"
            )
            or ""
        )

        external_document_id = str(
            document.get(
                "external_document_id"
            )
            or document.get(
                "cde_document_id"
            )
            or ""
        )

        if (
            not source_id
            or not external_document_id
        ):
            raise RuntimeError(
                (
                    "Customer-controlled document provenance "
                    "is incomplete."
                )
            )

        result = (
            self.db.table(
                "data_sources"
            )
            .select("*")
            .eq(
                "id",
                source_id,
            )
            .limit(1)
            .execute()
        )

        if not result.data:
            raise RuntimeError(
                "Linked CDE source was not found."
            )

        connector = ConnectorService.for_source(
            result.data[0]
        )

        raw = await connector.fetch_content(
            external_document_id
        )

        if not raw:
            raise RuntimeError(
                "The CDE returned an empty file."
            )

        return raw

    def _embed(
        self,
        chunks: list[str],
    ) -> list[list[float]]:
        vectors: list[
            list[float]
        ] = []

        for start in range(
            0,
            len(chunks),
            PROJECT_EMBED_BATCH,
        ):
            batch = chunks[
                start:
                start
                + PROJECT_EMBED_BATCH
            ]

            response = (
                self.openai.embeddings.create(
                    model=(
                        PROJECT_EMBEDDING_MODEL
                    ),
                    input=batch,
                )
            )

            ordered = sorted(
                response.data,
                key=lambda item:
                    item.index,
            )

            vectors.extend(
                [
                    list(
                        item.embedding
                    )
                    for item in ordered
                ]
            )

        if len(vectors) != len(
            chunks
        ):
            raise RuntimeError(
                (
                    "Embedding count did not match "
                    "document chunk count."
                )
            )

        return vectors

    def _ensure_collection(
        self,
        vector_size: int,
    ) -> None:
        from qdrant_client.models import (
            Distance,
            VectorParams,
        )

        try:
            info = self.qdrant.get_collection(
                PROJECT_QDRANT_COLLECTION
            )

            vectors = (
                info.config
                .params
                .vectors
            )

            configured_size = getattr(
                vectors,
                "size",
                None,
            )

            if (
                configured_size
                and int(
                    configured_size
                )
                != int(
                    vector_size
                )
            ):
                raise RuntimeError(
                    (
                        "Private Qdrant collection vector size "
                        f"is {configured_size}, but "
                        f"{PROJECT_EMBEDDING_MODEL} returned "
                        f"{vector_size}. Do not mix embedding models "
                        "inside the same collection."
                    )
                )

            return

        except RuntimeError:
            raise

        except Exception:
            pass

        self.qdrant.create_collection(
            collection_name=(
                PROJECT_QDRANT_COLLECTION
            ),
            vectors_config=VectorParams(
                size=vector_size,
                distance=Distance.COSINE,
            ),
        )

    def _delete_existing_vectors(
        self,
        *,
        tenant_id: str,
        project_id: str,
        document_id: str,
    ) -> None:
        from qdrant_client.models import (
            FieldCondition,
            Filter,
            MatchValue,
        )

        try:
            self.qdrant.delete(
                collection_name=(
                    PROJECT_QDRANT_COLLECTION
                ),
                points_selector=Filter(
                    must=[
                        FieldCondition(
                            key="tenant_id",
                            match=MatchValue(
                                value=tenant_id
                            ),
                        ),
                        FieldCondition(
                            key="project_id",
                            match=MatchValue(
                                value=project_id
                            ),
                        ),
                        FieldCondition(
                            key="document_id",
                            match=MatchValue(
                                value=document_id
                            ),
                        ),
                    ]
                ),
                wait=True,
            )

        except Exception:
            # A first-time document has nothing to delete.
            pass

    def _upsert_vectors(
        self,
        *,
        document: dict[str, Any],
        chunks: list[str],
        vectors: list[list[float]],
    ) -> None:
        from qdrant_client.models import (
            PointStruct,
        )

        tenant_id = str(
            document["tenant_id"]
        )

        project_id = str(
            document["project_id"]
        )

        document_id = str(
            document["id"]
        )

        file_hash = str(
            document.get(
                "content_hash"
            )
            or document.get(
                "file_hash"
            )
            or ""
        )

        self._delete_existing_vectors(
            tenant_id=tenant_id,
            project_id=project_id,
            document_id=document_id,
        )

        points: list[
            PointStruct
        ] = []

        for index, (
            chunk,
            vector,
        ) in enumerate(
            zip(
                chunks,
                vectors,
            )
        ):
            point_id = str(
                uuid.uuid5(
                    uuid.NAMESPACE_URL,
                    (
                        "contractx:"
                        f"{tenant_id}:"
                        f"{project_id}:"
                        f"{document_id}:"
                        f"{file_hash}:"
                        f"{index}"
                    ),
                )
            )

            payload = {
                "tenant_id":
                    tenant_id,
                "project_id":
                    project_id,
                "document_id":
                    document_id,
                "source_id":
                    document.get(
                        "source_id"
                    ),
                "source_type":
                    document.get(
                        "cde_source"
                    )
                    or document.get(
                        "storage_mode"
                    )
                    or "project_document",
                "storage_mode":
                    document.get(
                        "storage_mode"
                    ),
                "file_name":
                    document.get(
                        "file_name"
                    ),
                "document_name":
                    document.get(
                        "file_name"
                    ),
                "document_type":
                    document.get(
                        "document_category"
                    ),
                "category":
                    document.get(
                        "document_category"
                    ),
                "stakeholder_tier":
                    document.get(
                        "stakeholder_tier"
                    ),
                "flow_direction":
                    document.get(
                        "flow_direction"
                    ),
                "sender_entity":
                    document.get(
                        "sender_entity"
                    ),
                "recipient_entity":
                    document.get(
                        "recipient_entity"
                    ),
                "document_date":
                    document.get(
                        "date_sent"
                    ),
                "letter_reference_no":
                    document.get(
                        "letter_reference_no"
                    ),
                "subject_heading":
                    document.get(
                        "subject_heading"
                    ),
                "governing_clause_ref":
                    document.get(
                        "governing_clause_ref"
                    ),
                "is_formal_notice":
                    bool(
                        document.get(
                            "is_formal_notice"
                        )
                    ),
                "external_document_id":
                    document.get(
                        "external_document_id"
                    ),
                "external_path":
                    document.get(
                        "external_path"
                    ),
                "storage_path":
                    document.get(
                        "storage_path"
                    ),
                "file_hash":
                    document.get(
                        "file_hash"
                    ),
                "content_hash":
                    document.get(
                        "content_hash"
                    )
                    or document.get(
                        "file_hash"
                    ),
                "chunk_index":
                    index,
                "chunk_count":
                    len(chunks),
                "chunk_text":
                    chunk,
                "text":
                    chunk,
                "content":
                    chunk,
                "embedding_model":
                    PROJECT_EMBEDDING_MODEL,
                "indexed_at":
                    _utc_now(),
                "private_project_evidence":
                    True,
            }

            points.append(
                PointStruct(
                    id=point_id,
                    vector=vector,
                    payload=payload,
                )
            )

        self.qdrant.upsert(
            collection_name=(
                PROJECT_QDRANT_COLLECTION
            ),
            points=points,
            wait=True,
        )

    def _upsert_evidence_graph_node(
        self,
        *,
        document: dict[str, Any],
        chunk_count: int,
        character_count: int,
    ) -> None:
        tenant_id = str(
            document["tenant_id"]
        )

        project_id = str(
            document["project_id"]
        )

        document_id = str(
            document["id"]
        )

        existing = (
            self.db.table(
                "evidence_nodes"
            )
            .select("id")
            .eq(
                "tenant_id",
                tenant_id,
            )
            .eq(
                "project_id",
                project_id,
            )
            .eq(
                "node_type",
                "document",
            )
            .eq(
                "source_document_id",
                document_id,
            )
            .limit(1)
            .execute()
        )

        payload = {
            "document_category":
                document.get(
                    "document_category"
                ),
            "stakeholder_tier":
                document.get(
                    "stakeholder_tier"
                ),
            "flow_direction":
                document.get(
                    "flow_direction"
                ),
            "sender_entity":
                document.get(
                    "sender_entity"
                ),
            "recipient_entity":
                document.get(
                    "recipient_entity"
                ),
            "letter_reference_no":
                document.get(
                    "letter_reference_no"
                ),
            "governing_clause_ref":
                document.get(
                    "governing_clause_ref"
                ),
            "chunk_count":
                chunk_count,
            "character_count":
                character_count,
            "index_status":
                "ready",
        }

        provenance = {
            "source_id":
                document.get(
                    "source_id"
                ),
            "storage_mode":
                document.get(
                    "storage_mode"
                ),
            "storage_path":
                document.get(
                    "storage_path"
                ),
            "external_document_id":
                document.get(
                    "external_document_id"
                ),
            "external_path":
                document.get(
                    "external_path"
                ),
            "file_hash":
                document.get(
                    "file_hash"
                ),
            "content_hash":
                document.get(
                    "content_hash"
                )
                or document.get(
                    "file_hash"
                ),
            "qdrant_collection":
                PROJECT_QDRANT_COLLECTION,
            "embedding_model":
                PROJECT_EMBEDDING_MODEL,
            "indexed_at":
                _utc_now(),
        }

        values = {
            "tenant_id":
                tenant_id,
            "project_id":
                project_id,
            "node_type":
                "document",
            "source_document_id":
                document_id,
            "source_ref":
                document.get(
                    "letter_reference_no"
                ),
            "title":
                document.get(
                    "subject_heading"
                )
                or document.get(
                    "file_name"
                )
                or "Project Document",
            "event_date":
                document.get(
                    "date_sent"
                ),
            "payload":
                payload,
            "provenance":
                provenance,
            "confidence":
                1.0,
            "updated_at":
                _utc_now(),
        }

        if existing.data:
            (
                self.db.table(
                    "evidence_nodes"
                )
                .update(
                    values
                )
                .eq(
                    "id",
                    existing.data[0][
                        "id"
                    ],
                )
                .execute()
            )
        else:
            (
                self.db.table(
                    "evidence_nodes"
                )
                .insert(
                    values
                )
                .execute()
            )

    async def index_document(
        self,
        *,
        document_id: str,
        tenant_id: str,
        project_id: str,
        job_id: Optional[str] = None,
        raw: Optional[bytes] = None,
    ) -> dict[str, Any]:
        document = self._get_document(
            document_id=document_id,
            tenant_id=tenant_id,
            project_id=project_id,
        )

        self._update_document(
            document_id,
            {
                "index_status":
                    "processing",
                # UI intentionally remains at 10 until final success.
                "processing_percent":
                    10,
            },
        )

        self._update_job(
            job_id,
            {
                "stage":
                    "extracting",
                "processing_percent":
                    10,
                "current_task":
                    (
                        "Extracting and indexing "
                        "private project evidence"
                    ),
                "error_message":
                    None,
            },
        )

        try:
            file_bytes = (
                raw
                if raw is not None
                else await self._load_bytes(
                    document
                )
            )

            if not file_bytes:
                raise RuntimeError(
                    "Document bytes are empty."
                )

            actual_hash = hashlib.sha256(
                file_bytes
            ).hexdigest()

            expected_hash = str(
                document.get(
                    "content_hash"
                )
                or document.get(
                    "file_hash"
                )
                or ""
            )

            if (
                expected_hash
                and actual_hash
                != expected_hash
            ):
                raise RuntimeError(
                    (
                        "Document integrity check failed: "
                        "content hash does not match the "
                        "registered project document."
                    )
                )

            text = await asyncio.to_thread(
                extract_text,
                file_bytes,
                str(
                    document.get(
                        "file_name"
                    )
                    or ""
                ),
            )

            chunks = chunk_text(
                text
            )

            if not chunks:
                raise RuntimeError(
                    "No indexable chunks were produced."
                )

            vectors = await asyncio.to_thread(
                self._embed,
                chunks,
            )

            vector_size = len(
                vectors[0]
            )

            await asyncio.to_thread(
                self._ensure_collection,
                vector_size,
            )

            await asyncio.to_thread(
                self._upsert_vectors,
                document=document,
                chunks=chunks,
                vectors=vectors,
            )

            await asyncio.to_thread(
                self._upsert_evidence_graph_node,
                document=document,
                chunk_count=len(
                    chunks
                ),
                character_count=len(
                    text
                ),
            )

            self._update_document(
                document_id,
                {
                    "index_status":
                        "ready",
                    "processing_percent":
                        100,
                    "content_hash":
                        actual_hash,
                    "last_synced_at":
                        (
                            document.get(
                                "last_synced_at"
                            )
                            or _utc_now()
                        ),
                },
            )

            self._update_job(
                job_id,
                {
                    "stage":
                        "complete",
                    "transfer_percent":
                        100,
                    "processing_percent":
                        100,
                    "chunks_processed":
                        len(chunks),
                    "chunks_total":
                        len(chunks),
                    "current_task":
                        "Ready",
                    "completed_at":
                        _utc_now(),
                    "error_message":
                        None,
                },
            )

            return {
                "status":
                    "ready",
                "document_id":
                    document_id,
                "processing_percent":
                    100,
                "chunk_count":
                    len(chunks),
                "character_count":
                    len(text),
                "collection":
                    PROJECT_QDRANT_COLLECTION,
                "embedding_model":
                    PROJECT_EMBEDDING_MODEL,
                "evidence_graph_node":
                    "document",
            }

        except Exception as exc:
            self._update_document(
                document_id,
                {
                    "index_status":
                        "failed",
                    # Do not show a fake intermediate/final percentage.
                    "processing_percent":
                        10,
                },
            )

            self._update_job(
                job_id,
                {
                    "stage":
                        "failed",
                    "processing_percent":
                        10,
                    "current_task":
                        "Private indexing failed",
                    "error_message":
                        str(exc),
                    "completed_at":
                        _utc_now(),
                },
            )

            raise


_indexer: Optional[
    PrivateProjectIndexer
] = None


def get_private_project_indexer(
) -> PrivateProjectIndexer:
    global _indexer

    if _indexer is None:
        _indexer = (
            PrivateProjectIndexer()
        )

    return _indexer
