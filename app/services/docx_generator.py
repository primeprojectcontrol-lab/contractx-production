import io
from typing import Dict, Any, List
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex: str):
    """Sets background shading color for docx table cells."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def build_contract_review_docx(review_data: Dict[str, Any]) -> bytes:
    doc = docx.Document()
    
    # Configure 0.5-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Styles & Colors
    HEX_PRIMARY = "1E293B"      # Dark Slate
    HEX_TEXT_MUTED = "64748B"   # Muted Gray
    HEX_BG_LIGHT = "F8FAFC"     # Light Gray Background

    # Header Title
    title = doc.add_paragraph()
    run_title = title.add_run("ContractX.ai | Executive Contract Risk Audit")
    run_title.bold = True
    run_title.font.size = Pt(18)
    run_title.font.color.rgb = RGBColor.from_string(HEX_PRIMARY)

    meta = doc.add_paragraph()
    run_meta = meta.add_run(
        f"Project: {review_data.get('project_name', 'Default Project')} | "
        f"Document: {review_data.get('document_name', 'Sub-Contract Agreement')} | "
        f"Region: {review_data.get('region', 'UAE_GCC')}"
    )
    run_meta.font.size = Pt(9)
    run_meta.font.color.rgb = RGBColor.from_string(HEX_TEXT_MUTED)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Executive Summary Table
    summary_table = doc.add_table(rows=1, cols=1)
    summary_table.autofit = False
    summary_cell = summary_table.rows[0].cells[0]
    set_cell_background(summary_cell, HEX_BG_LIGHT)
    
    p = summary_cell.paragraphs[0]
    r_risk = p.add_run(f"Overall Risk Level: {review_data.get('overall_risk', 'HIGH').upper()}\n\n")
    r_risk.bold = True
    r_body = p.add_run(review_data.get('executive_summary', ''))
    r_body.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Risk & Deviations Table Header
    h2 = doc.add_paragraph()
    r_h2 = h2.add_run("Contractual Deviations & Risk Matrix")
    r_h2.bold = True
    r_h2.font.size = Pt(12)

    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    headers = ["Executed Clause", "Standard Reference", "Identified Deviation & Impact", "Risk Level"]
    col_widths = [Inches(1.2), Inches(1.5), Inches(3.8), Inches(1.0)]

    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], HEX_PRIMARY)
        for r in hdr_cells[i].paragraphs[0].runs:
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.bold = True
            r.font.size = Pt(9)

    # Fill Deviations Data
    deviations: List[Dict[str, Any]] = review_data.get("contractual_flaws_and_deviations", [])
    for dev in deviations:
        row_cells = table.add_row().cells
        row_cells[0].text = str(dev.get("executed_clause", "N/A"))
        row_cells[1].text = str(dev.get("standard_reference", "N/A"))
        row_cells[2].text = f"{dev.get('issue', '')}\n{dev.get('impact_analysis', '')}"
        row_cells[3].text = str(dev.get("risk_level", "MEDIUM")).upper()

        for idx, cell in enumerate(row_cells):
            cell.width = col_widths[idx]
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Strategic Advisory Section
    rec_h2 = doc.add_paragraph()
    r_rec_h2 = rec_h2.add_run("Strategic Legal Advisory & Risk Mitigation")
    r_rec_h2.bold = True
    r_rec_h2.font.size = Pt(12)

    rec_p = doc.add_paragraph()
    r_rec = rec_p.add_run(review_data.get("advisory_recommendation", ""))
    r_rec.font.size = Pt(9.5)

    # Export to Bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()